"""Excel parsing. Workbooks are treated as untrusted input; formulas are not executed."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from app.core.exceptions import PayloadTooLargeError, ValidationAppError
from app.utils.normalization import normalize_header, normalize_identifier

REQUIRED_HEADERS = {
    "BASE VEHICLE MODEL": "base_model_name",
    "VARIANT NAME": "variant_name",
    "RLF ID": "rlf_id",
    "RM ID": "rm_id",
    "IP ID": "ip_id",
    "EVAP ID": "evap_id",
    "PR ID": "pr_id",
    "DF ID": "df_id",
    "OB ID": "ob_id",
    "ER ID": "er_id",
    "PEMS ID": "pems_id",
}

ID_KEYS = ("rlf_id", "rm_id", "ip_id", "evap_id", "pr_id", "df_id", "ob_id", "er_id", "pems_id")


@dataclass
class ParsedVariant:
    name: str
    row: int


@dataclass
class ParsedVehicle:
    base_model_name: str
    identifiers: dict[str, str | None]
    variants: list[ParsedVariant] = field(default_factory=list)
    first_row: int = 0
    duplicate_variant_names: list[str] = field(default_factory=list)


@dataclass
class ParseIssue:
    row: int | None
    field: str | None
    message: str


@dataclass
class ParsedWorkbook:
    file_name: str
    total_rows: int
    vehicles: list[ParsedVehicle]
    errors: list[ParseIssue]
    warnings: list[ParseIssue]


def validate_upload(filename: str | None, content: bytes, max_bytes: int) -> None:
    if not filename:
        raise ValidationAppError("A file name is required", field="file")
    lowered = filename.lower()
    if not lowered.endswith(".xlsx"):
        raise ValidationAppError("Only .xlsx files are supported", field="file")
    if lowered.endswith(".xlsm") or lowered.endswith(".xls"):
        raise ValidationAppError("Macro-enabled and legacy Excel files are not supported", field="file")
    if len(content) > max_bytes:
        raise PayloadTooLargeError("File exceeds configured limit")
    if len(content) == 0:
        raise ValidationAppError("The uploaded file is empty", field="file")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_workbook(content: bytes, file_name: str) -> ParsedWorkbook:
    try:
        workbook: Workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_vba=False,
        )
    except Exception as exc:  # noqa: BLE001
        raise ValidationAppError("The uploaded file is not a valid .xlsx workbook", field="file") from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ValidationAppError("The workbook has no header row", field="file") from exc

        mapping = _map_headers(header_row)
        vehicles: list[ParsedVehicle] = []
        errors: list[ParseIssue] = []
        warnings: list[ParseIssue] = []
        grouped: dict[str, ParsedVehicle] = {}
        order: list[str] = []
        current_base: str | None = None
        data_row_count = 0

        for index, row in enumerate(rows, start=2):
            values = _row_values(row, mapping)
            if _is_empty_row(values):
                continue
            data_row_count += 1

            raw_base = values.get("base_model_name")
            if raw_base:
                current_base = raw_base
            elif not current_base:
                errors.append(
                    ParseIssue(
                        row=index,
                        field="base_model_name",
                        message="A variant appears before any base model has been defined",
                    )
                )
                continue

            variant_name = values.get("variant_name")
            if not variant_name:
                errors.append(
                    ParseIssue(row=index, field="variant_name", message="Variant name is required")
                )
                continue

            group_key = _model_key(current_base)
            if group_key not in grouped:
                identifiers = {key: values.get(key) for key in ID_KEYS}
                grouped[group_key] = ParsedVehicle(
                    base_model_name=current_base,
                    identifiers=identifiers,
                    first_row=index,
                )
                order.append(group_key)
            else:
                for key in ID_KEYS:
                    incoming = values.get(key)
                    existing = grouped[group_key].identifiers.get(key)
                    if incoming and existing is None:
                        grouped[group_key].identifiers[key] = incoming
                    elif incoming and existing and incoming.lower() != existing.lower():
                        warnings.append(
                            ParseIssue(
                                row=index,
                                field=key,
                                message=(
                                    f"Conflicting {key} for '{grouped[group_key].base_model_name}'; "
                                    f"keeping the first value from row {grouped[group_key].first_row}"
                                ),
                            )
                        )

            vehicle = grouped[group_key]
            existing = {item.name.lower() for item in vehicle.variants}
            if variant_name.lower() in existing:
                vehicle.duplicate_variant_names.append(variant_name)
                warnings.append(
                    ParseIssue(
                        row=index,
                        field="variant_name",
                        message=f"Duplicate variant '{variant_name}' under '{current_base}' will be skipped",
                    )
                )
                continue
            vehicle.variants.append(ParsedVariant(name=variant_name, row=index))

        vehicles = [grouped[name] for name in order]
        return ParsedWorkbook(
            file_name=file_name,
            total_rows=data_row_count,
            vehicles=vehicles,
            errors=errors,
            warnings=warnings,
        )
    finally:
        workbook.close()


def _map_headers(header_row: tuple[Any, ...]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        header = normalize_header(cell)
        if header in REQUIRED_HEADERS:
            found[REQUIRED_HEADERS[header]] = index

    missing = [label for label, key in REQUIRED_HEADERS.items() if key not in found]
    if missing:
        raise ValidationAppError(
            f"Required header(s) missing: {', '.join(missing)}",
            field="headers",
        )
    return found


def _model_key(value: str) -> str:
    return " ".join(value.casefold().split())


def _row_values(row: tuple[Any, ...], mapping: dict[str, int]) -> dict[str, str | None]:
    values: dict[str, str | None] = {}
    for key, index in mapping.items():
        raw = row[index] if index < len(row) else None
        if key in {"base_model_name", "variant_name"}:
            text = _cell_text(raw)
            values[key] = text or None
        else:
            values[key] = normalize_identifier(raw)
    return values


def _is_empty_row(values: dict[str, str | None]) -> bool:
    return all(not value for value in values.values())
